from __future__ import annotations

import csv
import re
import unicodedata
from collections import defaultdict
from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from conferencia.domain.box_codes import normalize_caixa_estoque


class ExcelReadError(Exception):
    pass


class OpenpyxlPalletReader:
    REQUIRED_COLUMN = "CAIXAESTOQUE"

    def read_carton_codes(self, file_path: Path, extension: str) -> list[str]:
        rows = self._xlsx_rows(file_path) if extension == ".xlsx" else self._csv_rows(file_path)
        try:
            header = next(rows)
        except StopIteration as error:
            raise ExcelReadError("O arquivo não possui cabeçalho.") from error
        index = self._required_column_index(header)
        codes: list[str] = []
        lines: dict[str, list[int]] = defaultdict(list)
        first_by_key: dict[str, str] = {}
        for line_number, row in enumerate(rows, start=2):
            if not any(self._raw_value(cell) for cell in row):
                continue
            if index >= len(row):
                continue
            code = self._cell_value(row[index], line_number)
            if code:
                key = code
                codes.append(code)
                lines[key].append(line_number)
                first_by_key.setdefault(key, code)
        if not codes:
            raise ExcelReadError("Não há caixas válidas na coluna CAIXA_ESTOQUE.")
        duplicated = [
            f"{first_by_key[key]} (linhas {', '.join(map(str, positions))})"
            for key, positions in lines.items()
            if len(positions) > 1
        ]
        if duplicated:
            raise ExcelReadError(
                "Caixas duplicadas no arquivo: " + "; ".join(duplicated[:10])
            )
        return codes

    def _xlsx_rows(self, path: Path) -> Iterator[list[object]]:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (InvalidFileException, OSError, ValueError) as error:
            raise ExcelReadError("Não foi possível ler o arquivo .xlsx.") from error
        try:
            for row in workbook.active.iter_rows():
                yield list(row)
        finally:
            workbook.close()

    def _csv_rows(self, path: Path) -> Iterator[list[str]]:
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as file:
                # O CSV operacional é separado por ponto e vírgula.  O módulo csv
                # entrega todas as células como texto, sem inferência de tipo.
                yield from csv.reader(file, delimiter=";")
        except (OSError, UnicodeError, csv.Error) as error:
            raise ExcelReadError("Não foi possível ler o arquivo .csv.") from error

    def _required_column_index(self, header: list[object]) -> int:
        for index, cell in enumerate(header):
            normalized = self._normalize_header(self._raw_value(cell))
            accepted = {
                self.REQUIRED_COLUMN,
                f"{self.REQUIRED_COLUMN}VARCHAR2",
                f"{self.REQUIRED_COLUMN}TEXT",
            }
            if normalized in accepted:
                return index
        received = [self._raw_value(cell) for cell in header]
        raise ExcelReadError(
            "A coluna obrigatória CAIXA_ESTOQUE não foi encontrada. "
            f"Cabeçalhos recebidos: {received!r}"
        )

    @staticmethod
    def _normalize_header(value: str) -> str:
        decomposed = unicodedata.normalize("NFKD", value)
        plain = "".join(
            character for character in decomposed if not unicodedata.combining(character)
        )
        return re.sub(r"[^A-Z0-9]", "", plain.upper())

    @staticmethod
    def _raw_value(cell: object) -> str:
        if cell is None:
            return ""
        value = cell.value if hasattr(cell, "value") else cell
        return normalize_caixa_estoque(value)

    @staticmethod
    def _cell_value(cell: object, line_number: int) -> str:
        value = cell.value if hasattr(cell, "value") else cell
        if value is None:
            return ""
        if isinstance(value, str):
            return normalize_caixa_estoque(value)
        number_format = getattr(cell, "number_format", "")
        if isinstance(value, int) and re.fullmatch(r"0+", number_format):
            return f"{value:0{len(number_format)}d}"
        raise ExcelReadError(
            f"Linha {line_number}: código numérico sem zeros preserváveis. "
            "Exporte CAIXA_ESTOQUE como texto pelo WMS."
        )
